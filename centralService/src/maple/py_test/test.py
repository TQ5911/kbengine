import json
import random
import sys
import time
import hmac
import hashlib
import requests
from openpyxl import load_workbook

BASE_URL = "http://192.168.10.133:9527"


def _auth_headers(method: str):
    ts = str(int(time.time()))
    nc = str(random.randint(100000, 999999))
    secret = "hello world"  # keep in sync with server config
    sign = hmac.new(secret.encode(), (method + ts + nc).encode(), hashlib.sha256).hexdigest()
    return {
        "Content-Type": "application/json",
        "ts": ts,
        "nc": nc,
        "sign": sign,
    }


def get_zone_dic_from_xlsx(wb):
    sheet = wb['分区表-zoneList']
    second_row = []
    for i in range(1, sheet.max_column + 1):
        if not sheet.cell(row=2, column=i).value:
            break

        second_row.append(sheet.cell(row=2, column=i).value)

    id_index = second_row.index("ID")
    zone_name_index = second_row.index("zoneName")

    zone_dic = {}
    for i in range(3, sheet.max_row + 1):
        zone_id = sheet.cell(row=i, column=id_index + 1).value
        zone_name = sheet.cell(row=i, column=zone_name_index + 1).value
        zone_dic[zone_id] = zone_name

    return zone_dic


def xlsx_to_json(file_name):
    wb = load_workbook(file_name)
    sheet = wb['服务器数据表-serverList']

    zone_dic = get_zone_dic_from_xlsx(wb)

    second_row = []
    for i in range(1, sheet.max_column + 1):
        if not sheet.cell(row=2, column=i).value:
            break

        second_row.append(sheet.cell(row=2, column=i).value)

    # ['ID', 'serverName', 'zoneID', 'zoneName', 'gameServer', 'queueServer', 'serverGroup', 'serverState', 'serverFlagState']
    id_index = second_row.index("ID")
    server_name_index = second_row.index("serverName")
    zone_id_index = second_row.index("zoneID")
    game_server_index = second_row.index("gameServer")
    queue_server_index = second_row.index("queueServer")
    server_group_index = second_row.index("serverGroup")
    server_state_index = second_row.index("serverState")
    server_flag_state_index = second_row.index("serverFlagState")

    _json_list = []
    for i in range(7, sheet.max_row + 1):
        _id = sheet.cell(row=i, column=id_index + 1).value
        server_name = sheet.cell(row=i, column=server_name_index + 1).value
        zone_id = sheet.cell(row=i, column=zone_id_index + 1).value
        zone_name = zone_dic[zone_id]
        game_server = sheet.cell(row=i, column=game_server_index + 1).value
        queue_server = sheet.cell(row=i, column=queue_server_index + 1).value
        server_group = sheet.cell(row=i, column=server_group_index + 1).value
        server_state = sheet.cell(row=i, column=server_state_index + 1).value
        server_flag_state = sheet.cell(row=i, column=server_flag_state_index + 1).value

        if not _id:
            break

        if isinstance(server_flag_state, int):
            server_flag_state = 1 << server_flag_state

        elif isinstance(server_flag_state, str):
            sps = server_flag_state.split(",")
            server_flag_state = 0
            for sp in sps:
                if sp:
                    server_flag_state |= 1 << int(sp)

        else:
            server_flag_state = 0

        _json_list.append({
            "id": int(_id),
            "server_name": server_name,
            "zone_id": zone_id,
            "zone_name": zone_name,
            "game_server": game_server,
            "queue_server": queue_server,
            "server_group": server_group,
            "server_state": server_state,
            "server_flag_state": server_flag_state,
        })

    return {
        "servers": _json_list
    }


def import_server_from_xlsx(file_name):
    payload = xlsx_to_json(file_name)
    url = f"{BASE_URL}/import"
    response = requests.post(url, headers=_auth_headers("POST"), json=payload)
    print(response.status_code, response.text)


def import_server_from_json(file_name):
    with open(file_name, "r", encoding="utf-8") as f:
        payload = json.load(f)
    url = f"{BASE_URL}/import"
    response = requests.post(url, headers=_auth_headers("POST"), json=payload)
    print(response.status_code, response.text)

def get_all_server():
    url = f"{BASE_URL}/getAllServer"
    response = requests.get(url)
    print(response.text)


def get_all_zone_data():
    url = f"{BASE_URL}/getAllZoneData"
    response = requests.get(url)
    print(response.text)


def add_zone():
    url = f"{BASE_URL}/addZone"
    ts = str(int(time.time()))
    nc = str(random.randint(100000, 999999))
    secret = "hello world"
    method = "POST"
    sign = hmac.new(secret.encode(), (method + ts + nc).encode(), hashlib.sha256).hexdigest()

    response = requests.post(
        url,
        json={"zone_id": 1, "zone_name": "test"},
        headers={
            "Content-Type": "application/json",
            "ts": ts,
            "nc": nc,
            "sign": sign})
    print(response.text)


def export_server_to_json(file_name):
    url = f"{BASE_URL}/export"
    response = requests.get(url, headers=_auth_headers("GET"))
    if response.status_code != 200:
        print(response.status_code, response.text)
        return
    data = response.json()
    with open(file_name, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"exported to {file_name}")


def _print_usage():
    print("Usage:")
    print("  python test.py import-xlsx <xlsx_file>")
    print("  python test.py import-json <json_file>")
    print("  python test.py export-json <out_json_file>")
    print("  python test.py list-servers")
    print("  python test.py list-zones")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        _print_usage()
        sys.exit(1)

    cmd = sys.argv[1]
    if cmd == "import-xlsx" and len(sys.argv) >= 3:
        import_server_from_xlsx(sys.argv[2])
    elif cmd == "import-json" and len(sys.argv) >= 3:
        import_server_from_json(sys.argv[2])
    elif cmd == "export-json" and len(sys.argv) >= 3:
        export_server_to_json(sys.argv[2])
    elif cmd == "list-servers":
        get_all_server()
    elif cmd == "list-zones":
        get_all_zone_data()
    else:
        _print_usage()
